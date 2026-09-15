"""Скрипт валидации отчета ЛР3 на соответствие ГОСТ 7.32."""
import docx

doc_path = r"c:\Users\Слава\Documents\university-labs\information-ecology\lab-03-noise-parameters\docs\Отчет_ЛР3_Смирнов_БСТ2556.docx"
doc = docx.Document(doc_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Проверка на паразитные \n
newline_errors = []
for i, p in enumerate(doc.paragraphs):
    if "\\n" in p.text:
        newline_errors.append(f"P[{i}]: {p.text}")

for t_idx, t in enumerate(doc.tables):
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            if "\\n" in cell.text:
                newline_errors.append(f"Table[{t_idx}] R{r_idx}C{c_idx}: {cell.text}")

print(f"Literal '\\n' errors count: {len(newline_errors)}")
if newline_errors:
    print("Errors found:", newline_errors)

# Проверка шрифтов
fonts = set()
sizes = set()
for p in doc.paragraphs:
    for r in p.runs:
        fonts.add(r.font.name)
        if r.font.size:
            sizes.add(r.font.size.pt)

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    fonts.add(r.font.name)
                    if r.font.size:
                        sizes.add(r.font.size.pt)

print(f"Fonts used: {fonts}")
print(f"Font sizes used (pt): {sorted(list(sizes))}")
print("Validation complete!")
