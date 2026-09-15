"""
Скрипт генерации отчета по ГОСТ 7.32-2017 для Лабораторной работы № 14р (№ 2).
Тема: «РАСЧЕТ ЗАЩИТНЫХ ПАРАМЕТРОВ ПРИ РАБОТЕ С СВЧ-ПЕРЕДАТЧИКОМ»
Дисциплина: «Информационная экология» («Безопасность жизнедеятельности»).
Студент: Смирнов Вячеслав Юрьевич, гр. БСТ2556, шифр ЗБСТ25066.
Вариант: 6 (Табл. 14.1 и 14.2).
"""

import os
import sys
import math
import docx
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Добавляем путь к модулю расчета
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
if CURRENT_DIR not in sys.path:
    sys.path.insert(0, CURRENT_DIR)

from microwave_shielding_calculator import (
    MicrowaveShieldingCalculator,
    get_variant_6_parameters,
    C_LIGHT,
    MU_0,
    WAVE_IMPEDANCE_0,
    N_ENERGY_DOSE_LIMIT
)

DOCS_DIR = os.path.abspath(os.path.join(CURRENT_DIR, "..", "docs"))
ASSETS_DIR = os.path.join(DOCS_DIR, "assets")
os.makedirs(DOCS_DIR, exist_ok=True)


def set_table_gost_borders(table):
    """Устанавливает тонкие черные рамки таблицы по ГОСТ 7.32."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_row_cant_split(row):
    """Предотвращает разрыв строки таблицы между страницами."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def add_p(doc, text="", space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          indent=1.25, bold=False, italic=False, size=14, line_spacing=1.5):
    """Создает абзац строго по ГОСТ 7.32."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(indent)
    if text:
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.bold = bold
        run.font.italic = italic
    return p


def add_section_heading(doc, text: str):
    """Заголовок структурного элемента по ГОСТ 7.32 (14 pt, Bold, без точки на конце)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_subsection_heading(doc, text: str):
    """Подзаголовок подраздела по ГОСТ 7.32."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_formula_block(doc, formula_text: str, num_str: str = ""):
    """
    Размещает формулу по центру, а номер формулы (1) строго по правому краю
    с помощью двухколоночной невидимой таблицы. Исключает разрывы и сползание.
    """
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(14.5)
    table.columns[1].width = Cm(2.0)

    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    c0 = table.rows[0].cells[0]
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.line_spacing = 1.5
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(4)
    p0.paragraph_format.first_line_indent = Cm(0)
    r0 = p0.add_run(formula_text)
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(14)
    r0.font.italic = True
    r0.font.color.rgb = RGBColor(0, 0, 0)

    c1 = table.rows[0].cells[1]
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.line_spacing = 1.5
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(4)
    p1.paragraph_format.first_line_indent = Cm(0)
    if num_str:
        r1 = p1.add_run(num_str)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(14)
        r1.font.color.rgb = RGBColor(0, 0, 0)


def build_gost_report():
    """Генерирует полный отчет по ЛР № 2 в строгом соответствии с ГОСТ 7.32-2017."""
    doc = Document()

    # 1. Настройка полей страницы
    section = doc.sections[0]
    section.left_margin = Cm(3.0)   # Левое поле 30 мм
    section.right_margin = Cm(1.5)  # Правое поле 15 мм
    section.top_margin = Cm(2.0)    # Верхнее поле 20 мм
    section.bottom_margin = Cm(2.0) # Нижнее поле 20 мм
    section.page_width = Cm(21.0)   # A4
    section.page_height = Cm(29.7)

    # 2. Вычисления параметров варианта 6
    gen_v6, mat_v6 = get_variant_6_parameters()
    res = MicrowaveShieldingCalculator.calculate(gen_v6, mat_v6)

    # ==============================================================================
    # ТИТУЛЬНЫЙ ЛИСТ (БЕЗ ПАРАЗИТНЫХ \n, АККУРАТНЫЕ ПАРАГРАФЫ И ТАБЛИЦА)
    # ==============================================================================
    add_p(doc, "ФЕДЕРАЛЬНОЕ АГЕНТСТВО СВЯЗИ",
          space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, size=12, line_spacing=1.15)
    add_p(doc, "Ордена Трудового Красного Знамени",
          space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=11, line_spacing=1.15)
    add_p(doc, "Федеральное государственное бюджетное образовательное учреждение высшего образования",
          space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=11, line_spacing=1.15)
    add_p(doc, "«МОСКОВСКИЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ СВЯЗИ И ИНФОРМАТИКИ»",
          space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, size=12, line_spacing=1.15)
    add_p(doc, "Кафедра экологии, безопасности жизнедеятельности и электропитания",
          space_before=0, space_after=48, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=12, line_spacing=1.15)

    add_p(doc, "ОТЧЕТ", space_before=18, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, size=18, line_spacing=1.15)
    add_p(doc, "по лабораторной работе № 14р (№ 2)", space_before=0, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, size=14, line_spacing=1.15)
    add_p(doc, "по дисциплине «Информационная экология»", space_before=0, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=14, line_spacing=1.15)
    add_p(doc, "(курс «Безопасность жизнедеятельности»)", space_before=0, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=12, italic=True, line_spacing=1.15)
    add_p(doc, "Тема: «РАСЧЕТ ЗАЩИТНЫХ ПАРАМЕТРОВ ПРИ РАБОТЕ С СВЧ-ПЕРЕДАТЧИКОМ»",
          space_before=0, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, size=14, line_spacing=1.15)
    add_p(doc, "(Вариант № 6)", space_before=0, space_after=42, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, bold=True, size=14, line_spacing=1.15)

    # Блок подписей (2-колоночная невидимая таблица)
    sign_table = doc.add_table(rows=1, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    sign_table.columns[0].width = Cm(7.5)
    sign_table.columns[1].width = Cm(9.0)

    tblPr = sign_table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    c_right = sign_table.rows[0].cells[1]

    def add_sign_line(cell, text, bold=False, space_after=2):
        if len(cell.paragraphs) == 1 and cell.paragraphs[0].text == "":
            p = cell.paragraphs[0]
        else:
            p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = bold
        r.font.color.rgb = RGBColor(0, 0, 0)

    add_sign_line(c_right, "Выполнил:")
    add_sign_line(c_right, "студент группы БСТ2556")
    add_sign_line(c_right, "Смирнов Вячеслав Юрьевич", bold=True)
    add_sign_line(c_right, "Номер зачетной книжки: ЗБСТ25066")
    add_sign_line(c_right, "Направление: 09.03.02 Информационные системы и технологии")
    add_sign_line(c_right, "Профиль: Безопасность компьютерных систем / DevSecOps (2 курс)", space_after=12)
    add_sign_line(c_right, "Проверил:")
    add_sign_line(c_right, "к.т.н., доцент кафедры ЭБЖиЭ")
    add_sign_line(c_right, "Шарофутдинов Р.М.", bold=True, space_after=0)

    add_p(doc, "Москва, 2026 г.", space_before=58, space_after=0, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=12, line_spacing=1.0)

    doc.add_page_break()

    # ==============================================================================
    # 1. ЦЕЛЬ РАБОТЫ И ЗАДАЧИ ИССЛЕДОВАНИЯ
    # ==============================================================================
    add_section_heading(doc, "1 ЦЕЛЬ РАБОТЫ И ЗАДАЧИ ИССЛЕДОВАНИЯ")

    add_p(doc,
          "Цель работы: исследование электромагнитной обстановки на рабочем месте оператора "
          "при эксплуатации СВЧ-генератора, освоение методики аналитической оценки плотности потока "
          "энергии излучения в волновой зоне и расчет параметров инженерно-технической защиты — "
          "минимальной толщины экранирующей камеры и длины волноводной трубки для вывода органа управления.")

    add_p(doc, "Задачи работы:")
    add_p(doc, "1) Рассчитать напряженность магнитной составляющей поля H катушки выходного контура на рабочем месте оператора без защитного экрана;")
    add_p(doc, "2) Проверить критерии формирования волновой (дальней) зоны электромагнитного излучения;")
    add_p(doc, "3) Определить фактическую плотность потока энергии (ППЭ) и сопоставить ее с нормативным предельно допустимым уровнем (ПДУ) облучения персонала;")
    add_p(doc, "4) Рассчитать требуемую кратность ослабления электромагнитного поля L;")
    add_p(doc, "5) Рассчитать минимальную теоретическую толщину металлического экрана delta_min и обосновать выбор конструктивной толщины;")
    add_p(doc, "6) Рассчитать погонное затухание в трубке-волноводе alpha и необходимую длину трубки l для экранированного вывода диэлектрического стержня ручки управления;")
    add_p(doc, "7) Сформулировать обоснованные инженерные выводы и ответить на контрольные вопросы.")

    # ==============================================================================
    # 2. ИСХОДНЫЕ ДАННЫЕ ВАРИАНТА № 6
    # ==============================================================================
    add_section_heading(doc, "2 ИСХОДНЫЕ ДАННЫЕ ВАРИАНТА № 6")

    add_p(doc,
          "В соответствии с шифром зачетной книжки ЗБСТ25066 вариант выбирается следующим образом: "
          "по последней цифре (6) — параметры выходного контура и геометрии рабочего места (Таблица 14.1), "
          "по предпоследней цифре (6) — характеристики конструкционных материалов экрана и стержня (Таблица 14.2).")

    add_p(doc, "Таблица 1 — Исходные данные для расчета (Вариант № 6)", bold=True, space_before=6, space_after=4, indent=0)

    table1 = doc.add_table(rows=1, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False
    set_table_gost_borders(table1)

    col_widths1 = [Cm(1.2), Cm(8.0), Cm(3.3), Cm(4.0)]
    for i, w in enumerate(col_widths1):
        table1.columns[i].width = w

    hdr_cells1 = table1.rows[0].cells
    hdr_titles1 = ["№", "Наименование параметра", "Обозначение", "Значение"]
    for i, t in enumerate(hdr_titles1):
        hdr_cells1[i].width = col_widths1[i]
        p = hdr_cells1[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(t)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True

    var_data1 = [
        ("1", "Число витков катушки индуктивности", "W", f"{gen_v6.W}"),
        ("2", "Сила тока в катушке", "I, А", f"{gen_v6.I:.1f}"),
        ("3", "Рабочая частота генератора", "f, Гц", "4 · 10⁸ (400 МГц)"),
        ("4", "Время облучения (регулировок) за смену", "Т, ч", f"{gen_v6.T:.1f} (12 мин)"),
        ("5", "Диаметр ручки управления (отверстия)", "D, м", "4 · 10⁻² (4 см)"),
        ("6", "Расстояние от катушки до рабочего места", "R, м", f"{gen_v6.R:.1f}"),
        ("7", "Радиус катушки индуктивности", "r, м", "10⁻¹ (0,10 м = 10 см)"),
        ("8", "Материал защитного экрана", "—", f"{mat_v6.screen_material}"),
        ("9", "Относительная магнитная проницаемость экрана", "μ", f"{mat_v6.mu:.1f}"),
        ("10", "Удельная электрическая проводимость экрана", "γ, См/м", "5,7 · 10⁷"),
        ("11", "Материал диэлектрического стержня ручки", "—", f"{mat_v6.rod_material}"),
        ("12", "Относительная диэлектрическая проницаемость стержня", "ε", f"{mat_v6.epsilon:.1f}"),
        ("13", "Нормативная предельная энергетическая нагрузка", "N, Вт·ч/м²", f"{N_ENERGY_DOSE_LIMIT:.1f}"),
    ]

    for row_data in var_data1:
        row = table1.add_row()
        set_row_cant_split(row)
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = col_widths1[i]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)

    # ==============================================================================
    # 3. ТЕОРЕТИЧЕСКИЕ ОСНОВЫ И РАСЧЕТНЫЕ СХЕМЫ
    # ==============================================================================
    add_section_heading(doc, "3 ТЕОРЕТИЧЕСКИЕ ОСНОВЫ И РАСЧЕТНЫЕ СХЕМЫ")

    add_p(doc,
          "В СВЧ-передатчиках источником интенсивного электромагнитного поля является выходной колебательный "
          "контур, содержащий катушку переменной индуктивности. Для обеспечения безопасных условий труда персонала "
          "контур помещается в экранирующую камеру. Однако ручка регулировки индуктивности должна выводиться наружу "
          "через стенку экрана, что создает канал утечки электромагнитной энергии.")

    add_p(doc,
          "На рисунке 1 приведена принципиальная схема взаимного расположения излучающего контура, защитного экрана "
          "с волноводной трубкой и рабочего места оператора.")

    # Вставка схемы 14.1
    fig1_path = os.path.join(ASSETS_DIR, "fig14_1_scheme.png")
    if os.path.exists(fig1_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.first_line_indent = Cm(0)
        run_img = p_img.add_run()
        run_img.add_picture(fig1_path, width=Cm(6.0))

    add_p(doc, "Рисунок 1 — Принципиальная схема экранирования СВЧ-генератора с выводом ручки управления",
          space_before=2, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=12, italic=True)

    add_p(doc,
          "Для исключения проникновения электромагнитного поля через отверстие в стенке экрана ручка управления "
          "выводится наружу при помощи полой металлической трубки длиной l, в которую помещается стержень из диэлектрика "
          "(эбонита). Такая конструкция (рисунок 2) представляет собой запредельный цилиндрический волновод, "
          "в котором электромагнитная волна рабочей частоты быстро затухает по экспоненциальному закону.")

    # Вставка схемы 14.2
    fig2_path = os.path.join(ASSETS_DIR, "fig14_2_handle.png")
    if os.path.exists(fig2_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(6)
        p_img2.paragraph_format.space_after = Pt(4)
        p_img2.paragraph_format.first_line_indent = Cm(0)
        run_img2 = p_img2.add_run()
        run_img2.add_picture(fig2_path, width=Cm(6.5))

    add_p(doc, "Рисунок 2 — Конструкция вывода ручки управления: 1 — диэлектрический стержень; 2 — металлическая трубка-волновод",
          space_before=2, space_after=12, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, size=12, italic=True)

    # ==============================================================================
    # 4. ПОШАГОВЫЕ ЭЛЕКТРОДИНАМИЧЕСКИЕ РАСЧЕТЫ
    # ==============================================================================
    add_section_heading(doc, "4 ПОШАГОВЫЕ ЭЛЕКТРОДИНАМИЧЕСКИЕ РАСЧЕТЫ")

    # 4.1 Магнитное поле
    add_subsection_heading(doc, "4.1 Расчет напряженности магнитной составляющей поля катушки")
    add_p(doc,
          "Напряженность магнитной составляющей поля H на расстоянии R от катушки индуктивности (без экрана) "
          "определяется формулой:")
    add_formula_block(doc, "H = (W · I · r²) / (4 · R³) · β_m, (А/м)", "(1)")
    add_p(doc,
          f"где β_m — поправочный коэффициент, зависящий от соотношения расстояния R и радиуса катушки r. "
          f"В рассматриваемом случае R / r = {gen_v6.R:.1f} / {gen_v6.r:.1f} = {res.ratio_R_r:.1f}. "
          f"Так как R / r > 10, коэффициент принимается равным β_m = 1.")
    add_p(doc, "Подставляя исходные данные варианта № 6, получаем:")
    add_formula_block(doc, f"H = (12 · 80 · 0,10²) / (4 · 2,0³) · 1 = 9,6 / 32 = {res.H:.4f} А/м.", "(2)")

    # 4.2 Волновая зона
    add_subsection_heading(doc, "4.2 Проверка условий волновой (дальней) зоны излучения")
    add_p(doc,
          "Оценка воздействия электромагнитного поля по плотности потока энергии (ППЭ) корректна только в том случае, "
          "если рабочее место оператора находится в волновой зоне (зоне излучения). Длина электромагнитной волны lambda "
          "при частоте генератора f = 4 · 10⁸ Гц равна:")
    add_formula_block(doc, f"λ = c / f = (3 · 10⁸) / (4 · 10⁸) = {res.wavelength:.2f} м.", "(3)")
    add_p(doc, "Условия формирования волновой зоны определяются неравенствами:")
    add_formula_block(doc, "R >> λ / (2 · π)    и    R >> r² / λ.", "(4)")
    add_p(doc, "Вычислим граничные пороговые значения:")
    add_formula_block(doc, f"λ / (2 · π) = {res.wavelength:.2f} / (2 · 3,1416) = {res.condition_1_limit:.4f} м ≈ 0,12 м;", "(5)")
    add_formula_block(doc, f"r² / λ = 0,10² / {res.wavelength:.2f} = 0,01 / 0,75 = {res.condition_2_limit:.4f} м ≈ 0,013 м.", "(6)")
    add_p(doc,
          f"Фактическое расстояние до рабочего места составляет R = {gen_v6.R:.1f} м, что превышает пороговое значение "
          f"λ / (2π) в {gen_v6.R / res.condition_1_limit:.1f} раза, а пороговое значение r² / λ — в {gen_v6.R / res.condition_2_limit:.1f} раз. "
          f"Следовательно, оба условия R >> λ/(2π) и R >> r²/λ строго выполняются, и на рабочем месте оператора "
          f"действительно имеет место волновая зона излучения.")

    # 4.3 ППЭ и ПДУ
    add_subsection_heading(doc, "4.3 Оценка интенсивности излучения без экрана и нормативного ПДУ")
    add_p(doc,
          "В волновой зоне оценка интенсивности электромагнитного поля производится по плотности потока энергии (ППЭ), "
          "рассчитываемой по соотношению Умова–Пойнтинга через волновое сопротивление свободного пространства (377 Ом):")
    add_formula_block(doc, "Π = 377 · H² / 2, (Вт/м²).", "(7)")
    add_formula_block(doc, f"Π = (377 · {res.H:.4f}²) / 2 = (377 · 0,09) / 2 = 33,93 / 2 = {res.ppe:.4f} Вт/м².", "(8)")

    add_p(doc,
          "Предельно допустимая плотность потока энергии (ППЭ_доп) регламентируется санитарными нормами исходя из "
          "допустимой интегральной энергетической нагрузки за рабочую смену N = 2 Вт·ч/м² и суммарного времени "
          f"воздействия излучения T = {gen_v6.T:.1f} ч (12 минут):")
    add_formula_block(doc, "Π_доп = N / T, (Вт/м²).", "(9)")
    add_formula_block(doc, f"Π_доп = 2,0 / {gen_v6.T:.1f} = {res.ppe_dop:.1f} Вт/м².", "(10)")

    add_p(doc,
          f"Сопоставление фактической плотности потока энергии с нормативным значением показывает: "
          f"Π = {res.ppe:.4f} Вт/м² > Π_доп = {res.ppe_dop:.1f} Вт/м². "
          f"На рабочем месте наблюдается превышение допустимого уровня облучения в {res.hazard_ratio:.2f} раза. "
          f"Следовательно, эксплуатация генератора без защитного экранирования недопустима, и требуется установка "
          f"экранирующей камеры.")

    # 4.4 Ослабление L
    add_subsection_heading(doc, "4.4 Расчет требуемого ослабления электромагнитного поля")
    add_p(doc,
          "Требуемый коэффициент ослабления электромагнитного поля L, который должен обеспечивать экран, равен:")
    add_formula_block(doc, "L = Π / Π_доп.", "(11)")
    add_formula_block(doc, f"L = {res.ppe:.4f} / {res.ppe_dop:.1f} = {res.L:.4f}.", "(12)")
    add_p(doc, "В логарифмических единицах (децибелах) требуемая величина затухания составляет:")
    add_formula_block(doc, f"L_дБ = 10 · lg(L) = 10 · lg({res.L:.4f}) = 10 · {math.log10(res.L):.4f} = {res.L_db:.2f} дБ.", "(13)")

    # 4.5 Толщина экрана
    add_subsection_heading(doc, "4.5 Расчет минимальной толщины металлического экрана")
    add_p(doc,
          "Ослабление электромагнитной волны в проводящем слое металла обусловлено поверхностным эффектом (скин-эффектом) "
          "и вихревыми токами Фуко. Зная электрофизические характеристики материала экрана (медь: μ = 1, γ = 5,7 · 10⁷ См/м), "
          "рассчитаем круговую частоту ω и абсолютную магнитную проницаемость μ_a:")
    add_formula_block(doc, f"ω = 2 · π · f = 2 · 3,1416 · (4 · 10⁸) = {res.omega:.4e} рад/с;", "(14)")
    add_formula_block(doc, f"μ_a = μ_0 · μ = (4 · π · 10⁻⁷) · 1 = {res.mu_a:.4e} Гн/м.", "(15)")

    add_p(doc, "Глубина проникновения электромагнитного поля в медь (толщина скин-слоя Delta) на частоте 400 МГц равна:")
    add_formula_block(doc, "Δ = √(2 / (ω · μ_a · γ)).", "(16)")
    add_formula_block(doc, f"Δ = √(2 / ({res.omega:.3e} · {res.mu_a:.3e} · 5,7·10⁷)) = {res.skin_depth:.4e} м = {res.skin_depth*1e6:.2f} мкм.", "(17)")

    add_p(doc,
          "Минимальная расчетная толщина стенки экрана delta, обеспечивающая заданное ослабление L, определяется по формуле методички:")
    add_formula_block(doc, "δ = ln(L) / (2 · √(ω · μ_a · γ / 2)).", "(18)")
    add_p(doc,
          f"Заметим, что знаменатель формулы (18) равен 2 / Δ. Следовательно, расчетная формула может быть представлена как "
          f"δ = (Δ · ln(L)) / 2. Подставляя значения ln({res.L:.4f}) = {math.log(res.L):.4f} и Δ = {res.skin_depth*1e6:.2f} мкм:")
    add_formula_block(doc, f"δ_min = ({res.skin_depth*1e6:.2f} · 10⁻⁶ · {math.log(res.L):.4f}) / 2 = {res.delta_min:.4e} м = {res.delta_min_um:.3f} мкм ({res.delta_min_mm:.6f} мм).", "(19)")
    add_p(doc,
          f"Теоретически требуемая электродинамическая толщина медного экрана чрезвычайно мала (менее 1 мкм), "
          f"что объясняется высокой частотой поля (400 МГц) и превосходной электрической проводимостью меди. "
          f"Однако из соображений механической прочности, жесткости конструкции и устойчивости к деформациям "
          f"конструктивная толщина медного экрана принимается равной δ_констр = {res.delta_constructive_mm:.1f} мм "
          f"(что с колоссальным запасом перекрывает расчетную величину).")

    # 4.6 Трубка-волновод
    add_subsection_heading(doc, "4.6 Расчет параметров трубки-волновода для вывода ручки управления")
    add_p(doc,
          "Металлическая трубка круглого сечения диаметром D = 4 · 10⁻² м (4 см), в которую помещен диэлектрический стержень "
          "из эбонита (ε = 3,0), работает в режиме запредельного волновода. Критическая длина волны основного типа колебаний "
          "H_11 в круглом волноводе с диэлектрическим заполнением составляет:")
    lambda_cr = 1.706 * gen_v6.D * math.sqrt(mat_v6.epsilon)
    add_formula_block(doc, f"λ_кр = 1,706 · D · √ε = 1,706 · 0,04 · √3,0 = {lambda_cr:.4f} м.", "(20)")
    add_p(doc,
          f"Поскольку длина рабочей волны λ = {res.wavelength:.2f} м значительно больше критической (λ = 0,75 м > λ_кр = {lambda_cr:.2f} м), "
          f"волна не может распространяться внутри трубки и экспоненциально затухает. "
          f"Ослабление электромагнитной энергии в трубке-волноводе на 1 метр длины определяется формулой:")
    add_formula_block(doc, "α = 32 / (D · √ε), (дБ/м).", "(21)")
    add_formula_block(doc, f"α = 32 / (0,04 · √3,0) = 32 / (0,04 · 1,7321) = 32 / 0,06928 = {res.alpha:.2f} дБ/м.", "(22)")

    add_p(doc,
          f"Необходимая расчетная длина трубки l_min, обеспечивающая требуемое ослабление поля L_дБ = {res.L_db:.2f} дБ, "
          f"составляет:")
    add_formula_block(doc, "l = (10 · lg(L)) / α, (м).", "(23)")
    add_formula_block(doc, f"l_min = {res.L_db:.2f} / {res.alpha:.2f} = {res.l_tube_min:.5f} м = {res.l_tube_min_mm:.2f} мм ≈ 5 мм.", "(24)")
    add_p(doc,
          f"Минимальная расчетная длина трубки составляет {res.l_tube_min_mm:.2f} мм. "
          f"С конструктивной точки зрения, а также для обеспечения надежного механического крепления в стенке камеры "
          f"и исключения краевых эффектов рассеяния, длина металлической трубки выбирается не менее ее диаметра (D = 40 мм) "
          f"и принимается равной l_констр = {res.l_tube_constructive_mm:.0f} мм.")

    # ==============================================================================
    # 5. СВОДНАЯ ТАБЛИЦА РЕЗУЛЬТАТОВ
    # ==============================================================================
    add_section_heading(doc, "5 СВОДНАЯ ТАБЛИЦА РЕЗУЛЬТАТОВ РАСЧЕТА")

    add_p(doc, "Таблица 2 — Итоговые параметры электродинамического расчета экранирования", bold=True, space_before=6, space_after=4, indent=0)

    table2 = doc.add_table(rows=1, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    set_table_gost_borders(table2)

    col_widths2 = [Cm(1.2), Cm(8.5), Cm(3.0), Cm(3.8)]
    for i, w in enumerate(col_widths2):
        table2.columns[i].width = w

    hdr_cells2 = table2.rows[0].cells
    hdr_titles2 = ["№", "Параметр", "Формула", "Значение"]
    for i, t in enumerate(hdr_titles2):
        hdr_cells2[i].width = col_widths2[i]
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(t)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True

    results_data = [
        ("1", "Напряженность магнитного поля катушки без экрана", "H = (W·I·r²)/(4·R³)·β_m", f"{res.H:.4f} А/м"),
        ("2", "Длина электромагнитной волны", "λ = c / f", f"{res.wavelength:.2f} м"),
        ("3", "Пороговые критерии волновой зоны", "λ/(2π) ; r²/λ", f"{res.condition_1_limit:.3f} м ; {res.condition_2_limit:.3f} м"),
        ("4", "Фактическая плотность потока энергии без экрана", "Π = 377·H² / 2", f"{res.ppe:.4f} Вт/м²"),
        ("5", "Предельно допустимая ППЭ за время смены T = 0,2 ч", "Π_доп = N / T", f"{res.ppe_dop:.2f} Вт/м²"),
        ("6", "Кратность превышения санитарной нормы", "Π / Π_доп", f"в {res.hazard_ratio:.2f} раза"),
        ("7", "Требуемый коэффициент ослабления поля", "L = Π / Π_доп", f"{res.L:.4f} ({res.L_db:.2f} дБ)"),
        ("8", "Глубина проникновения поля в медь (скин-слой)", "Δ = √(2/(ω·μ_a·γ))", f"{res.skin_depth*1e6:.2f} мкм"),
        ("9", "Минимальная расчетная толщина медного экрана", "δ = ln(L)/(2√(ωμ_aγ/2))", f"{res.delta_min_um:.3f} мкм"),
        ("10", "Рекомендуемая конструктивная толщина экрана", "δ_констр", f"{res.delta_constructive_mm:.1f} мм"),
        ("11", "Погонное затухание в трубке-волноводе", "α = 32 / (D·√ε)", f"{res.alpha:.2f} дБ/м"),
        ("12", "Минимальная расчетная длина трубки", "l = 10·lg(L) / α", f"{res.l_tube_min_mm:.2f} мм"),
        ("13", "Рекомендуемая конструктивная длина трубки", "l_констр", f"{res.l_tube_constructive_mm:.0f} мм"),
    ]

    for row_data in results_data:
        row = table2.add_row()
        set_row_cant_split(row)
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = col_widths2[i]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)

    # ==============================================================================
    # 6. ВЫВОДЫ
    # ==============================================================================
    add_section_heading(doc, "6 ВЫВОДЫ")

    add_p(doc,
          f"1. В результате проведенного электродинамического анализа установлено, что при работе СВЧ-генератора "
          f"на частоте f = 400 МГц на расстоянии R = 2,0 м формируется волновая зона электромагнитного поля "
          f"(условия R >> λ/(2π) = 0,12 м и R >> r²/λ = 0,013 м строго выполняются). Напряженность магнитного поля "
          f"без экранирования составляет H = {res.H:.3f} А/м, а фактическая плотность потока энергии достигает "
          f"Π = {res.ppe:.3f} Вт/м².")

    add_p(doc,
          f"2. При регламентированном времени регулировок T = 0,2 ч (12 минут в смену) предельно допустимый уровень "
          f"облучения составляет Π_доп = {res.ppe_dop:.1f} Вт/м². Фактическая плотность потока энергии превышает ПДУ "
          f"в {res.hazard_ratio:.2f} раза, что создает угрозу здоровью оператора и требует обязательного экранирования.")

    add_p(doc,
          f"3. Для снижения интенсивности излучения до безопасных нормативных значений требуется коэффициент ослабления "
          f"L = {res.L:.3f} (или L_дБ = {res.L_db:.2f} дБ). Расчетная минимальная толщина медного экрана составляет "
          f"δ_min = {res.delta_min_um:.3f} мкм, что обусловлено высокой проводимостью меди (γ = 5,7 · 10⁷ См/м) "
          f"и малым скин-слоем (Δ = {res.skin_depth*1e6:.2f} мкм). По условиям механической прочности и устойчивости к деформациям "
          f"толщина защитного кожуха принимается конструктивно равной δ_констр = {res.delta_constructive_mm:.1f} мм.")

    add_p(doc,
          f"4. Вывод ручки управления с эбонитовым стержнем (ε = 3,0) диаметром D = 4 см через стенку экрана обеспечивает "
          f"погонное затухание запредельного волновода α = {res.alpha:.2f} дБ/м. Минимальная теоретическая длина трубки "
          f"составляет l_min = {res.l_tube_min_mm:.2f} мм. Для надежной конструктивной фиксации в стенке экранирующей камеры "
          f"и гарантированного подавления излучения конструктивная длина трубки принимается равной l_констр = {res.l_tube_constructive_mm:.0f} мм.")

    # ==============================================================================
    # 7. ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ
    # ==============================================================================
    add_section_heading(doc, "7 ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ")

    add_subsection_heading(doc, "Вопрос 1. По каким параметрам оценивается степень облучения персонала, обслуживающего СВЧ-генераторы?")
    add_p(doc,
          "Ответ: В диапазоне сверхвысоких частот (СВЧ: от 300 МГц до 300 ГГц, длины волн от 1 м до 1 мм) электромагнитное "
          "поле в рабочей зоне персонала практически всегда формируется в дальней (волновой) зоне излучения, где волновое "
          "сопротивление среды стабилизируется на значении 377 Ом, а электрическая и магнитная составляющие синфазны и "
          "связаны жестким соотношением E / H = 377. В связи с этим оценка степени облучения персонала производится "
          "по следующим базовым параметрам:")
    add_p(doc,
          "1) Плотность потока энергии (ППЭ, вектор Пойнтинга Π) — характеризует поверхностную мощность излучения, падающую "
          "на единицу облучаемой площади тела человека за единицу времени. Единицы измерения: Вт/м² или мкВт/см² "
          "(1 Вт/м² = 100 мкВт/см²);")
    add_p(doc,
          "2) Энергетическая нагрузка (ЭН_Π, доза облучения) — интегральная величина поглощенной энергии за рабочую смену, "
          "учитывающая кумулятивный эффект воздействия: ЭН_Π = Π · T (Вт·ч/м² или мкВт·ч/см²), где T — продолжительность "
          "воздействия электромагнитного поля в часах;")
    add_p(doc,
          "3) Частота излучения f (или длина волны λ) — определяет глубину проникновения поля в биологические ткани человека "
          "и степень теплового/нетермогенного поглощения;")
    add_p(doc,
          "4) Временной режим и пространственное распределение — непрерывный, импульсный или сканирующий (прерывистый) "
          "характер облучения, скважность импульсов, пиковая мощность.")

    add_subsection_heading(doc, "Вопрос 2. Каков ПДУ облучения при работе в СВЧ-диапазоне?")
    add_p(doc,
          "Ответ: В соответствии с нормативными требованиями СанПиН 1.2.3685-21 «Гигиенические нормативы и требования "
          "к обеспечению безопасности и (или) безвредности для человека факторов среды обитания» и ГОСТ 12.1.006-84 "
          "«ССБТ. Электромагнитные поля радиочастот. Допустимые уровни на рабочих местах»:")
    add_p(doc,
          "1) Нормирование осуществляется по предельно допустимой энергетической нагрузке за рабочий день (смену), "
          "которая для персонала составляет N = 2,0 Вт·ч/м² (200 мкВт·ч/см²);")
    add_p(doc,
          "2) Предельно допустимая плотность потока энергии рассчитывается с учетом времени пребывания в рабочей зоне: "
          "Π_доп = N / T;")
    add_p(doc,
          "3) При непрерывном облучении в течение полного 8-часового рабочего дня (T = 8 ч) ПДУ составляет "
          "Π_доп = 2,0 / 8 = 0,25 Вт/м² = 25 мкВт/см²;")
    add_p(doc,
          "4) При кратковременном воздействии (например, регулировка аппаратуры в течение T = 0,2 ч = 12 минут, как в работе) "
          "допустимый уровень возрастает до Π_доп = 2,0 / 0,2 = 10 Вт/м² (1000 мкВт/см²);")
    add_p(doc,
          "5) Санитарными нормами установлен безусловный верхний предел: независимо от малой продолжительности работы "
          "плотность потока энергии на рабочем месте оператора не должна превышать 10 Вт/м² (1000 мкВт/см²). Пребывание "
          "в зонах с уровнями выше 10 Вт/м² без специальных средств индивидуальной защиты категорически запрещено.")

    add_subsection_heading(doc, "Вопрос 3. Технические и индивидуальные средства защиты от электромагнитных излучений.")
    add_p(doc,
          "Ответ: Комплексная система защиты персонала от электромагнитных полей радиочастот и СВЧ-диапазона включает "
          "организационные, инженерно-технические и индивидуальные средства:")
    add_p(doc,
          "А. Организационные и планировочные методы:")
    add_p(doc,
          "— Защита расстоянием: удаление постоянных рабочих мест от СВЧ-установок, поскольку интенсивность поля убывает "
          "пропорционально квадрату (или кубу) расстояния;")
    add_p(doc,
          "— Защита временем: строгий хронометраж и ограничение времени нахождения персонала в зоне излучения;")
    add_p(doc,
          "— Рациональное размещение технологического оборудования, направленности антенн и секторов излучения.")
    add_p(doc,
          "Б. Коллективные инженерно-технические средства защиты:")
    add_p(doc,
          "1) Экранирование источников излучения: сплошные металлические кожухи (медь, латунь, алюминий, сталь), сетчатые экраны "
          "с размером ячейки намного меньше длины волны, а также многослойные и комбинированные экраны;")
    add_p(doc,
          "2) Радиопоглощающие покрытия и материалы: ферритовые пластины, графитонаполненные полимеры, эластичные резиновые коврики "
          "и волокнистые поглотители, трансформирующие энергию поля в теплоту;")
    add_p(doc,
          "3) Защита технологических отверстий и выводов: применение запредельных круглых и прямоугольных волноводов (металлических трубок), "
          "в которых волны рабочей частоты экспоненциально затухают; установка сотовых решеток и фильтров радиопомех на линиях питания;")
    add_p(doc,
          "4) Эквиваленты нагрузки и поглотители мощности: использование водяных или коаксиальных согласованных нагрузок при регулировке "
          "и калибровке передатчиков вместо открытого излучения в цех/лабораторию.")
    add_p(doc,
          "В. Средства индивидуальной защиты (СИЗ):")
    add_p(doc,
          "1) Защитная спецодежда: комбинезоны, халаты, полукомбинезоны и капюшоны из специальной металлизированной ткани "
          "(с микропроводом из латуни или нержавеющей стали, ткани с напылением серебра), обеспечивающие экранирование на 20–30 дБ;")
    add_p(doc,
          "2) Защитные очки: очки закрытого типа со стеклами, покрытыми тончайшим токопроводящим оптически прозрачным слоем диоксида олова "
          "(тип ОРЗ-5), предохраняющие хрусталик глаза от коагуляции белка и катаракты;")
    add_p(doc,
          "3) Защитная обувь и перчатки из электропроводящих эластомеров, исключающие контактные токи.")

    # ==============================================================================
    # 8. СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ
    # ==============================================================================
    add_section_heading(doc, "8 СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")

    sources = [
        "Курбатов В.А. Защита от СВЧ-излучения: Лабораторная работа № 14р по курсу «Безопасность жизнедеятельности» / В.А. Курбатов. — М.: МТУСИ, 2019. — 12 с.",
        "ГОСТ 7.32-2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления. — М.: Стандартинформ, 2018. — 28 с.",
        "ГОСТ 12.1.006-84. Система стандартов безопасности труда. Электромагнитные поля радиочастот. Допустимые уровни на рабочих местах и требования к проведению контроля. — М.: Издательство стандартов, 1985. — 18 с.",
        "СанПиН 1.2.3685-21. Гигиенические нормативы и требования к обеспечению безопасности и (или) безвредности для человека факторов среды обитания. — Введ. 01.03.2021. — М.: Роспотребнадзор, 2021.",
        "Безопасность жизнедеятельности: Учебник для вузов / Под ред. С.В. Белова. — 8-е изд., перераб. и доп. — М.: Высшая школа, 2009. — 616 с.",
        "Кузнецов В.Д. Электромагнитная экология и электромагнитная безопасность: Учебное пособие. — М.: Горячая линия - Телеком, 2015. — 240 с."
    ]

    for i, s in enumerate(sources, 1):
        add_p(doc, f"{i}. {s}")

    # Сохраняем файл
    target_path = os.path.join(DOCS_DIR, "Отчет_ЛР2_Смирнов_БСТ2556.docx")
    doc.save(target_path)
    print(f"Отчет успешно сформирован и сохранен: {target_path}")
    return target_path


if __name__ == "__main__":
    build_gost_report()
