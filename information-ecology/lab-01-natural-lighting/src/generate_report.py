import os
DOCS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'docs'))
os.makedirs(DOCS_DIR, exist_ok=True)
"""
Скрипт формирования полноценного отчета по лабораторной работе №1
в формате Microsoft Word (.docx) в строгом соответствии с ГОСТ 7.32
и методическими требованиями кафедры ЭБЖиЭ МТУСИ.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from lighting_calculator import NaturalLightingCalculator, get_variant_6_params


def set_cell_background(cell, hex_color: str):
    """Установка фонового цвета ячейки таблицы."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Установка тонких аккуратных границ для таблицы."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.15, first_indent=0.49):
    """Настройка стандартного абзаца."""
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if first_indent > 0:
        p.paragraph_format.first_line_indent = Inches(first_indent)


def add_heading_styled(doc, text: str, level: int = 1):
    """Добавление стилизованного заголовка."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=12, space_after=6, line_spacing=1.15, first_indent=0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(24, 43, 73)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(40, 70, 110)
    else:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(60, 60, 60)
    return p


def build_docx_report():
    params = get_variant_6_params()
    res = NaturalLightingCalculator.calculate(params)

    doc = Document()

    # Поля страницы по ГОСТ: левое 30 мм, правое 15 мм, верх/низ 20 мм
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.79)     # 20 mm
        section.bottom_margin = Inches(0.79)  # 20 mm
        section.left_margin = Inches(1.18)    # 30 mm
        section.right_margin = Inches(0.59)   # 15 mm

    # ==============================================================================
    # ТИТУЛЬНЫЙ ЛИСТ (МТУСИ)
    # ==============================================================================
    p_min = doc.add_paragraph()
    p_min.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_min, space_before=0, space_after=2, line_spacing=1.0, first_indent=0)
    r_min = p_min.add_run("МИНИСТЕРСТВО ЦИФРОВОГО РАЗВИТИЯ, СВЯЗИ И МАССОВЫХ КОММУНИКАЦИЙ\nРОССИЙСКОЙ ФЕДЕРАЦИИ")
    r_min.font.name = "Times New Roman"
    r_min.font.size = Pt(10)
    r_min.font.bold = True

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_univ, space_before=2, space_after=18, line_spacing=1.0, first_indent=0)
    r_univ = p_univ.add_run("Ордена Трудового Красного Знамени федеральное государственное бюджетное\n"
                           "образовательное учреждение высшего образования\n"
                           "«МОСКОВСКИЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ СВЯЗИ И ИНФОРМАТИКИ»\n(МТУСИ)")
    r_univ.font.name = "Times New Roman"
    r_univ.font.size = Pt(11)
    r_univ.font.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_sub, space_before=0, space_after=36, line_spacing=1.0, first_indent=0)
    r_sub = p_sub.add_run("Центр заочного обучения по программам бакалавриата\n"
                          "Кафедра «Экология, безопасность жизнедеятельности и электропитание»")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_title, space_before=24, space_after=6, line_spacing=1.15, first_indent=0)
    r_title = p_title.add_run("ОТЧЕТ\nПО ЛАБОРАТОРНОЙ РАБОТЕ № 1")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(24, 43, 73)

    p_disc = doc.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_disc, space_before=0, space_after=12, line_spacing=1.15, first_indent=0)
    r_disc = p_disc.add_run("по дисциплине «Информационная экология»")
    r_disc.font.name = "Times New Roman"
    r_disc.font.size = Pt(13)
    r_disc.font.italic = True

    p_topic = doc.add_paragraph()
    p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_topic, space_before=6, space_after=60, line_spacing=1.15, first_indent=0)
    r_topic = p_topic.add_run("Тема: «Расчёт естественной освещённости в производственном помещении»\n(Вариант № 6)")
    r_topic.font.name = "Times New Roman"
    r_topic.font.size = Pt(14)
    r_topic.font.bold = True

    # Блок автора и преподавателя
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    format_paragraph(p_author, space_before=12, space_after=60, line_spacing=1.15, first_indent=0)
    r_auth = p_author.add_run(
        "Выполнил: студент группы БСТ2556\n"
        "Смирнов Вячеслав Юрьевич\n"
        "Зачетная книжка: ЗБСТ25066\n"
        "Направление: 09.03.02 Информационные системы и технологии\n"
        "Профиль: Инженерия DevSecOps (2 курс)\n\n"
        "Проверил: доцент кафедры ЭБЖиЭ\n"
        "Курбатов В.А.\n"
    )
    r_auth.font.name = "Times New Roman"
    r_auth.font.size = Pt(12)

    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_city, space_before=36, space_after=0, line_spacing=1.0, first_indent=0)
    r_city = p_city.add_run("Москва, 2026 г.")
    r_city.font.name = "Times New Roman"
    r_city.font.size = Pt(12)

    doc.add_page_break()

    # ==============================================================================
    # 1. ЦЕЛЬ И ЗАДАЧИ РАБОТЫ
    # ==============================================================================
    add_heading_styled(doc, "1. Цель и задачи работы", level=1)

    p_goal = doc.add_paragraph()
    format_paragraph(p_goal)
    p_goal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_goal.add_run(
        "Цель работы: изучить методику нормирования и инженерного расчёта естественного "
        "бокового освещения в производственных помещениях вычислительных центров и аппаратных "
        "залов предприятий связи в соответствии с требованиями СП 52.13330 (СНиП 23-05-95) "
        "и методическими указаниями кафедры ЭБЖиЭ МТУСИ."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p_tasks = doc.add_paragraph()
    format_paragraph(p_tasks)
    p_tasks.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_tasks.add_run(
        "Задачи работы:\n"
        "1. На основе нормативного разряда зрительной работы определить нормированное значение "
        "коэффициента естественной освещённости (КЕО l_min).\n"
        "2. Вычислить геометрические параметры помещения, световых проемов и взаимного расположения "
        "с противостоящим зданием (h1, A/B, B/h1, L/H).\n"
        "3. Определить нормативные коэффициенты: световую характеристику окна (η0), коэффициент "
        "затенения противостоящими зданиями (K_зд), коэффициент светопропускания (r0) и средневзвешенный "
        "коэффициент отражения внутренних поверхностей (ρ_ср).\n"
        "4. Рассчитать необходимую площадь световых проемов S0 для обеспечения нормируемой освещенности.\n"
        "5. Определить необходимое число окон n и величину межоконных промежутков b, провести инженерный "
        "анализ равномерности освещения и реализуемости односторонней и двусторонней схем размещения."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # ==============================================================================
    # 2. ИСХОДНЫЕ ДАННЫЕ ВАРИАНТА №6
    # ==============================================================================
    add_heading_styled(doc, "2. Исходные данные варианта № 6", level=1)

    p_init_intro = doc.add_paragraph()
    format_paragraph(p_init_intro)
    p_init_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_init_intro.add_run(
        "Исходные данные приняты по Таблице 4.8 (последняя цифра зачетной книжки 6) "
        "и Таблице 4.9 (предпоследняя цифра зачетной книжки 6) методического пособия МТУСИ. "
        "Объект проектирования: аппаратный зал телеграфа / вычислительный зал на 3 этаже узла связи."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # Таблица исходных данных
    table_init = doc.add_table(rows=12, cols=3)
    table_init.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_init)

    init_data = [
        ("Разряд зрительной работы", "III", "Высокая точность (размер объекта различения 0,3–0,5 мм)"),
        ("Длина помещения (A)", "20,0 м", "Размер вдоль наружной светонесущей стены"),
        ("Ширина помещения (B)", "10,0 м", "Глубина помещения"),
        ("Высота помещения (h)", "5,0 м", "Высота от пола до потолка"),
        ("Высота оконного проема (h0)", "3,5 м", "Вертикальный размер окна"),
        ("Высота подоконника над полом (h_под)", "1,2 м", "Расстояние от пола до подоконника"),
        ("Высота рабочей поверхности (h_раб)", "1,2 м", "Уровень горизонтальной рабочей поверхности"),
        ("Расстояние до здания напротив (L)", "30,0 м", "Расстояние между противостоящими зданиями"),
        ("Высота карниза противостоящего здания (H)", "30,0 м", "Высота над уровнем подоконника зала"),
        ("Ширина одного окна (b0)", "2,0 м", "Горизонтальный размер светового проема"),
        ("Коэффициенты отражения поверхностей", "ρ_п = 50%\nρ_ст = 50%\nρ_пол = 10%", "Потолок и стены светлые, пол темный линолеум/бетон"),
        ("Тип остекления и категория помещения", "Категория Б\n(стальные переплеты, двойные)", "Аппаратный зал с незначительным выделением пыли")
    ]

    hdr_cells = table_init.rows[0].cells
    hdr_cells[0].text = "Наименование параметра"
    hdr_cells[1].text = "Значение"
    hdr_cells[2].text = "Примечание"
    for cell in hdr_cells:
        set_cell_background(cell, "E8EEF5")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.bold = True
                run.font.size = Pt(11)

    for idx, (param, val, note) in enumerate(init_data):
        row_cells = table_init.rows[idx].cells
        row_cells[0].text = param
        row_cells[1].text = val
        row_cells[2].text = note
        for c_idx, cell in enumerate(row_cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

    # ==============================================================================
    # 3. МЕТОДИКА И ПОШАГОВЫЙ РАСЧЕТ
    # ==============================================================================
    add_heading_styled(doc, "3. Методика и расчет параметров естественного освещения", level=1)

    p_f1 = doc.add_paragraph()
    format_paragraph(p_f1)
    p_f1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_f1.add_run(
        "3.1. Определение расчетного параметра окна h1\n"
        "Под параметром окна h1 понимается возвышение верхнего края оконного проема над "
        "горизонтальной рабочей поверхностью:\n"
        "h1 = h0 + h_под - h_раб\n"
        f"h1 = 3,5 + 1,2 - 1,2 = {res.h1:.2f} м."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p_f2 = doc.add_paragraph()
    format_paragraph(p_f2)
    p_f2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_f2.add_run(
        "3.2. Определение геометрических соотношений помещения\n"
        f"• Отношение длины помещения к ширине: A / B = 20,0 / 10,0 = {res.ratio_AB:.2f}\n"
        f"• Отношение ширины к расчетной высоте окна: B / h1 = 10,0 / {res.h1:.2f} = {res.ratio_Bh1:.3f}\n"
        f"• Отношение расстояния между зданиями к высоте карниза: L / H = 30,0 / 30,0 = {res.ratio_LH:.2f}"
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p_f3 = doc.add_paragraph()
    format_paragraph(p_f3)
    p_f3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_f3.add_run(
        "3.3. Выбор нормативных коэффициентов по таблицам пособия\n"
        f"1. Нормированное значение КЕО l_min (Таблица 4.1): для III разряда зрительной работы "
        f"при боковом освещении l_min = {res.l_min:.1f}%.\n"
        f"2. Коэффициент затемнения противостоящим зданием K_зд (Таблица 4.3): при L / H = 1,0 "
        f"коэффициент составляет K_зд = {res.K_zd:.2f}.\n"
        f"3. Световая характеристика окна η0 (Таблица 4.2): для отношения A / B = 2,0 и B / h1 = 2,857 "
        f"путем интерполяции между значениями для B/h1=2,5 (η0=13,0) и B/h1=3,0 (η0=18,0) получаем:\n"
        f"η0 = 13,0 + (18,0 - 13,0) × (2,857 - 2,5) / (3,0 - 2,5) = {res.eta0:.2f}.\n"
        f"4. Общий коэффициент светопропускания r0 (Таблица 4.5): для помещений категории Б "
        f"при вертикальном остеклении и стальных двойных переплетах r0 = {res.r0:.2f}."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p_f4 = doc.add_paragraph()
    format_paragraph(p_f4)
    p_f4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_f4.add_run(
        "3.4. Определение средневзвешенного коэффициента отражения ρ_ср и коэффициента r1\n"
        "Площади ограждающих поверхностей производственного помещения:\n"
        f"• Площадь пола: S_пол = A × B = 20 × 10 = {res.S_floor:.1f} м²;\n"
        f"• Площадь потолка: S_п = A × B = 20 × 10 = {res.S_ceil:.1f} м²;\n"
        f"• Площадь стен: S_ст = 2 × (A + B) × h = 2 × (20 + 10) × 5,0 = {res.S_walls:.1f} м²;\n"
        f"• Полная внутренняя площадь: S_общ = {res.S_floor:.1f} + {res.S_ceil:.1f} + {res.S_walls:.1f} = {res.S_total:.1f} м².\n"
        f"Средневзвешенный коэффициент отражения внутренних поверхностей:\n"
        f"ρ_ср = (ρ_п·S_п + ρ_ст·S_ст + ρ_пол·S_пол) / S_общ\n"
        f"ρ_ср = (0,50 × 200 + 0,50 × 300 + 0,10 × 200) / 700 = (100 + 150 + 20) / 700 = 270 / 700 = {res.rho_avg:.4f} ({res.rho_avg*100:.1f}%).\n"
        f"По Таблице 4.6 определяем коэффициент r1, учитывающий отраженный свет:\n"
        f"• При одностороннем боковом освещении: r1 = {res.r1_single:.2f};\n"
        f"• При двустороннем боковом освещении: r1 = {res.r1_double:.2f}."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p_f5 = doc.add_paragraph()
    format_paragraph(p_f5)
    p_f5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_f5.add_run(
        "3.5. Расчет площади световых проемов S0 и числа окон для одностороннего освещения\n"
        "Необходимая площадь световых проемов определяется по основной формуле:\n"
        "S0 = (S_пол · l_min · η0 · K_зд) / (100 · r0 · r1)\n"
        f"S0 = ({res.S_floor:.1f} × {res.l_min:.1f} × {res.eta0:.2f} × {res.K_zd:.2f}) / (100 × {res.r0:.2f} × {res.r1_single:.2f}) = {res.S0_single:.2f} м².\n"
        f"Площадь одного типового окна: S_окна = h0 × b0 = 3,5 × 2,0 = {res.S_window:.2f} м².\n"
        f"Расчетное число окон: n_точн = S0 / S_окна = {res.S0_single:.2f} / {res.S_window:.2f} = {res.n_exact_single:.2f} шт.\n"
        f"Принимаем ближайшее целое число окон: n = {res.n_single} шт.\n"
        f"Величина межоконного расстояния b при размещении n окон в наружной стене длиной A=20 м:\n"
        f"b = (A - n · b0) / (n + 1) = (20,0 - {res.n_single} × 2,0) / ({res.n_single} + 1) = (20,0 - {res.n_single * params.b0:.1f}) / {res.n_single + 1} = {res.b_single:.2f} м."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # ==============================================================================
    # 4. ИНЖЕНЕРНЫЙ АНАЛИЗ И ВЫБОР СХЕМЫ РАЗМЕЩЕНИЯ ОКОН
    # ==============================================================================
    add_heading_styled(doc, "4. Инженерный анализ и выбор рациональной схемы освещения", level=1)

    p_an1 = doc.add_paragraph()
    format_paragraph(p_an1)
    p_an1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_an1.add_run(
        "Анализ односторонней схемы освещения:\n"
        f"Принятое расчетное число окон n = {res.n_single} шт. при ширине каждого окна b0 = 2,0 м "
        f"требует суммарной ширины остекления {res.n_single * params.b0:.1f} м, что превышает общую "
        f"длину наружной стены помещения A = 20,0 м. Значение межоконного промежутка получилось отрицательным "
        f"(b = {res.b_single:.2f} м). Это означает, что физически разместить необходимое количество световых "
        f"проемов в один ряд по одной продольной стене НЕВОЗМОЖНО."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    p_an2 = doc.add_paragraph()
    format_paragraph(p_an2)
    p_an2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_an2.add_run(
        "Инженерное решение — переход на двустороннее боковое освещение:\n"
        "Для обеспечения нормативной освещенности глубокого зала (B = 10 м) и равномерного распределения "
        "светового потока целесообразно организовать световые проемы в двух противоположных наружных стенах. "
        "При двустороннем освещении нормативный коэффициент r1 по Таблице 4.6 снижается до r1 = 1,70, "
        "что отражает более глубокое взаимное проникновение света с обеих сторон:\n"
        f"S0_двуст = ({res.S_floor:.1f} × {res.l_min:.1f} × {res.eta0:.2f} × {res.K_zd:.2f}) / (100 × {res.r0:.2f} × {res.r1_double:.2f}) = {res.S0_double:.2f} м².\n"
        f"Расчетное число окон для двух стен: n_точн = {res.S0_double:.2f} / 7,0 = {res.n_exact_double:.2f} шт.\n"
        f"Принимаем четное число окон для симметричного размещения: n = {res.n_double} шт. (по {res.n_per_wall} окон на каждой стене).\n"
        f"Межоконное расстояние на каждой из двух стен длиной 20 м составит:\n"
        f"b = (A - {res.n_per_wall} × b0) / ({res.n_per_wall} + 1) = (20,0 - {res.n_per_wall * params.b0:.1f}) / {res.n_per_wall + 1} = {res.b_double:.2f} м.\n"
        f"Величина межоконного простенка b = {res.b_double:.2f} м обеспечивает высокую конструктивную надежность "
        f"несущих простенков здания и равномерность естественного освещения рабочего пространства."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # Сводная таблица сравнения схем
    table_comp = doc.add_table(rows=6, cols=3)
    table_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_comp)

    comp_headers = table_comp.rows[0].cells
    comp_headers[0].text = "Показатель"
    comp_headers[1].text = "Одностороннее освещение"
    comp_headers[2].text = "Двустороннее освещение (рекомендовано)"
    for cell in comp_headers:
        set_cell_background(cell, "E8EEF5")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.bold = True
                run.font.size = Pt(11)

    comp_rows = [
        ("Коэффициент отраженного света r1", f"{res.r1_single:.2f}", f"{res.r1_double:.2f}"),
        ("Требуемая площадь остекления S0", f"{res.S0_single:.2f} м²", f"{res.S0_double:.2f} м²"),
        ("Принятое число окон n", f"{res.n_single} шт. (на 1 стену)", f"{res.n_double} шт. (по {res.n_per_wall} шт. на 2 стены)"),
        ("Межоконный промежуток b", f"{res.b_single:.2f} м (не реализуемо)", f"{res.b_double:.2f} м (оптимально)"),
        ("Равномерность освещенности", "Низкая (затухание света к дальней стене)", "Высокая (перекрестное освещение)")
    ]

    for idx, (param, val1, val2) in enumerate(comp_rows):
        row_cells = table_comp.rows[idx + 1].cells
        row_cells[0].text = param
        row_cells[1].text = val1
        row_cells[2].text = val2
        for c_idx, cell in enumerate(row_cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

    # ==============================================================================
    # 5. ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ
    # ==============================================================================
    add_heading_styled(doc, "5. Ответы на контрольные вопросы", level=1)

    questions = [
        (
            "Вопрос 1: Как нормируется естественная освещенность?",
            "Ответ: Естественная освещенность нормируется в относительных единицах — через коэффициент "
            "естественной освещенности (КЕО, e_N, выражается в процентах). КЕО представляет собой отношение "
            "естественной освещенности, создаваемой в некоторой точке заданной плоскости внутри помещения "
            "светом неба (непосредственно или после отражений), к одновременному значению наружной горизонтальной "
            "освещенности, создаваемой светом полностью открытого небосвода: e = (E_внут / E_нар) × 100%. "
            "Нормирование в относительных величинах обусловлено непрерывным изменением наружной естественной "
            "освещенности в течение суток и сезонов года в зависимости от положения солнца и облачности."
        ),
        (
            "Вопрос 2: Какие виды естественного освещения вы знаете?",
            "Ответ: В соответствии со строительными нормами различают три основных вида естественного освещения:\n"
            "1. Боковое освещение — осуществляется через световые проемы (окна) в наружных стенах здания. "
            "Может быть односторонним (окна на одной стене) или двусторонним (окна на противоположных или смежных стенах).\n"
            "2. Верхнее освещение — осуществляется через световые аэрационные или зенитные фонари, проемы в перекрытиях "
            "и светопрозрачные конструкции кровли.\n"
            "3. Комбинированное (смешанное) освещение — сочетание бокового и верхнего естественного освещения, "
            "обеспечивающее наиболее высокую освещенность и равномерность в крупногабаритных производственных цехах."
        ),
        (
            "Вопрос 3: Какие требования предъявляются к системам производственного освещения?",
            "Ответ: К производственному освещению предъявляются следующие обязательные санитарно-гигиенические "
            "и технико-экономические требования:\n"
            "• Обеспечение нормативного уровня освещенности на рабочих поверхностях в соответствии с разрядом зрительной работы;\n"
            "• Равномерное распределение яркости в поле зрения и исключение резких теней на рабочих местах;\n"
            "• Отсутствие прямой и отраженной слепящей блесткости (ограничение показателя ослепленности);\n"
            "• Постоянство освещенности во времени (отсутствие стробоскопического эффекта и пульсаций светового потока);\n"
            "• Обеспечение правильной цветопередачи и спектрального состава света, близкого к естественному;\n"
            "• Пожаро- и электробезопасность осветительных установок, надежность, удобство эксплуатации и энергоэффективность."
        ),
        (
            "Вопрос 4: По каким параметрам определяется разряд зрительной работы?",
            "Ответ: Разряд зрительной работы определяется по СП 52.13330 (СНиП 23-05-95) исходя из следующих параметров:\n"
            "1. Наименьший эквивалентный размер объекта различения (в миллиметрах). Например: до 0,15 мм — разряд I (наивысшая точность); "
            "от 0,15 до 0,3 мм — разряд II (очень высокая точность); от 0,3 до 0,5 мм — разряд III (высокая точность, как в варианте 6); "
            "от 0,5 до 1 мм — разряд IV (средняя точность); более 5 мм — разряд VI (грубая работа).\n"
            "2. Контраст объекта различения с фоном (малый, средний, большой), определяемый соотношением яркостей объекта и фона.\n"
            "3. Характеристика фона (светлый, средний, темный) в зависимости от коэффициента отражения поверхности фона."
        ),
        (
            "Вопрос 5: Что характеризует «спектральная видность», в чем заключается ее особенность? Энергетические и фотометрические величины.",
            "Ответ: Спектральная чувствительность глаза (относительная спектральная световая эффективность, или спектральная видность V(λ)) "
            "характеризует избирательную чувствительность человеческого зрительного анализатора к электромагнитному излучению "
            "различных длин волн оптического диапазона при одинаковой мощности лучистого потока.\n"
            "Особенность: максимум чувствительности глаза в условиях дневного (колбочкового) зрения приходится на длину волны λ = 555 нм "
            "(желто-зеленая область спектра). При сумеречном (палочковом) зрении максимум сдвигается в сине-зеленую область (λ = 507 нм, эффект Пуркинье).\n"
            "Связь величин:\n"
            "• Энергетические величины характеризуют излучение безотносительно зрительного восприятия человека в абсолютных энергетических единицах "
            "(поток излучения Ф_е в ваттах Вт, энергетическая светимость Вт/м², энергетическая яркость Вт/(ср·м²)).\n"
            "• Фотометрические (световые) величины оценивают воздействие оптического излучения на орган зрения с учетом спектральной чувствительности V(λ): "
            "световой поток Ф измеряется в люменах (лм), освещенность E — в люксах (лк = лм/м²), сила света I — в канделах (кд), яркость L — в кд/м²."
        )
    ]

    for q, ans in questions:
        p_q = doc.add_paragraph()
        format_paragraph(p_q, space_before=6, space_after=2, first_indent=0)
        rq = p_q.add_run(q)
        rq.font.name = "Times New Roman"
        rq.font.bold = True
        rq.font.size = Pt(12)
        rq.font.color.rgb = RGBColor(24, 43, 73)

        p_a = doc.add_paragraph()
        format_paragraph(p_a, space_before=0, space_after=8, first_indent=0.49)
        p_a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ra = p_a.add_run(ans)
        ra.font.name = "Times New Roman"
        ra.font.size = Pt(11)

    # ==============================================================================
    # 6. ВЫВОДЫ
    # ==============================================================================
    add_heading_styled(doc, "6. Выводы", level=1)

    p_concl = doc.add_paragraph()
    format_paragraph(p_concl)
    p_concl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_concl.add_run(
        f"1. В ходе выполнения лабораторной работы № 1 освоена методика расчета систем естественного освещения "
        f"производственных помещений связи и IT-инфраструктуры по СП 52.13330 (СНиП 23-05-95).\n"
        f"2. Для аппаратного зала телеграфа размерами 20 × 10 × 5 м (III разряд зрительной работы, КЕО l_min = 2,0%) "
        f"выполнен полный пошаговый расчет световых параметров с учетом затенения противостоящим зданием (K_зд = {res.K_zd:.2f}), "
        f"светопропускания окон (r0 = {res.r0:.2f}) и отражающих свойств поверхностей (ρ_ср = {res.rho_avg*100:.1f}%, r1 = {res.r1_single:.2f}).\n"
        f"3. Расчет показал, что для одностороннего бокового освещения требуемая площадь остекления составляет S0 = {res.S0_single:.2f} м², "
        f"что соответствует n = {res.n_single} окнам шириной по 2,0 м. Суммарная ширина оконных проемов ({res.n_single * params.b0:.1f} м) "
        f"превышает длину стены здания (20,0 м), что делает одностороннюю схему физически нереализуемой (b = {res.b_single:.2f} м < 0).\n"
        f"4. Предложено рациональное инженерное решение: перейти на двустороннее боковое освещение. "
        f"При двусторонней схеме (r1 = {res.r1_double:.2f}) требуемая площадь окон составляет S0 = {res.S0_double:.2f} м², "
        f"что реализуется установкой {res.n_double} окон (по {res.n_per_wall} окон на каждой продольной стене). "
        f"Межоконный промежуток составляет b = {res.b_double:.2f} м, что обеспечивает конструктивную прочность простенков, "
        f"высокую равномерность естественной освещенности рабочих мест и соблюдение требований охраны труда."
    )
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # ==============================================================================
    # 7. СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ
    # ==============================================================================
    add_heading_styled(doc, "7. Список использованных источников", level=1)

    sources = [
        "1. Курбатов В.А. Расчёт естественной освещённости в производственном помещении: Учебно-методическое пособие по курсу «Безопасность жизнедеятельности» / В.А. Курбатов. — М.: МТУСИ, 2022. — 24 с.",
        "2. СП 52.13330.2016. Естественное и искусственное освещение. Актуализированная редакция СНиП 23-05-95*. — М.: Минстрой России, 2016.",
        "3. СанПиН 1.2.3685-21. Гигиенические нормативы и требования к обеспечению безопасности и (или) безвредности для человека факторов среды обитания.",
        "4. ГОСТ 12.1.046-2014. Система стандартов безопасности труда. Строительство. Нормы освещения строительных площадок.",
        "5. ГОСТ 7.32-2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления."
    ]

    for src in sources:
        p_s = doc.add_paragraph()
        format_paragraph(p_s, space_before=2, space_after=4, first_indent=0.49)
        p_s.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rs = p_s.add_run(src)
        rs.font.name = "Times New Roman"
        rs.font.size = Pt(11)

    target_docx = os.path.join(DOCS_DIR, "Отчет_ЛР1_Смирнов_БСТ2556.docx")
    doc.save(target_docx)
    print(f"Report successfully generated at: {target_docx}")


if __name__ == "__main__":
    build_docx_report()
